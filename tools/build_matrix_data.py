"""Genera matrix-data.js a partir de las bibliotecas MATRIX.

Uso:  python3 tools/build_matrix_data.py <Biblioteca_Maestra.xlsx> <Si_1.docx> <Interpretaciones_770_v2.xlsx> \
              <Tabla_Programas_Colas_y_Zonas.xlsx> [salida.js]

Fuente prioritaria: Biblioteca Maestra v1.x (arcanos, 32 puntos, contextos, zonas,
programas publicables, 26 colas). El .docx aporta el diccionario semántico
Arcano × tipo de zona y los núcleos de programas. Los registros en cuarentena
(Documentado_Publicable = NO) nunca se exportan.

v2 (Interpretaciones Arcano × Posición): las 770 lecturas se exportan como 22 arcanos base + 35 posiciones
+ plantillas; el script verifica que las plantillas reproducen EXACTAMENTE los 770 × 6 textos del Excel.
Para el modo Hombre se añade la versión masculina de cada campo que la necesite (tools/genero.py).
"""
import json, re, sys, collections, os
import openpyxl, docx
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import genero

MAESTRA, DOCX, V2, TABLA = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
OUT = sys.argv[5] if len(sys.argv) > 5 else "matrix-data.js"

wb = openpyxl.load_workbook(MAESTRA, read_only=True, data_only=True)
def sheet(name):
    rows = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if r[0] is not None]
def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""

# ---- 22 Arcanos canónicos ----
ARC = {}
for a in sheet("01_ARCANOS"):
    ARC[int(a["Arcano_ID"])] = dict(n=clean(a["Nombre"]), rec=clean(a["Recurso_Integrado"]),
        hipo=clean(a["Hipo"]), hiper=clean(a["Hiper"]), clave=clean(a["Clave_Integracion"]))
assert len(ARC) == 22

# ---- 32 puntos ----
PTS = {}
for p in sheet("02_PUNTOS"):
    PTS[clean(p["Punto"])] = dict(zona=clean(p["Zona_Principal"]), func=clean(p["Funcion_Canonica"]),
        geo=clean(p["Geometria"]), formula=clean(p["Formula"]))
assert len(PTS) == 32

# ---- 35 contextos: se derivan plantillas y se verifica que reproducen los 770 párrafos ----
INT = sheet("05_INTERPRETACIONES")
CTX = collections.OrderedDict()
for r in INT:
    cid = clean(r["Contexto_ID"])
    if cid in CTX: continue
    a = ARC[int(r["Arcano_ID"])]
    t = clean(r["Parrafo_Editorial"])
    for key in ("rec", "hipo", "hiper", "clave", "n"):     # los campos largos primero
        t = t.replace(a[key], "{" + key + "}")
    CTX[cid] = dict(punto=clean(r["Punto"]), zona=clean(r["Zona_ID"]), func=clean(r["Funcion"]), tpl=t)
def render(tpl, a): return re.sub(r"\{(\w+)\}", lambda m: a[m.group(1)], tpl)
bad = [r["ID"] for r in INT if render(CTX[clean(r["Contexto_ID"])]["tpl"], ARC[int(r["Arcano_ID"])]) != clean(r["Parrafo_Editorial"])]
assert not bad, f"plantillas no reproducen: {bad[:5]}"

# ---- Zonas y secuencias de activación (orden importa) ----
ZON = {}
for z in sheet("03_ZONAS"):
    ZON[clean(z["Zona_ID"])] = dict(nombre=clean(z["Nombre"]), puntos=clean(z["Puntos"]), lectura=clean(z["Lectura"]))
SEQ = [dict(id=clean(z["Zona_ID"]), nombre=clean(z["Zona"]), pts=[clean(z["Punto_1"]), clean(z["Punto_2"]), clean(z["Punto_3"])])
       for z in sheet("09_REGLAS_ACTIVACION")]

# ---- Lente de zona para programas (se extrae de las expresiones programa×zona) ----
LENS = {}
for z in sheet("08_PROGRAMA_ZONA"):
    m = re.search(r"se interpreta aquí en relación con (.+?), sin trasladar|se interpreta en relación con (.+?)\. Como", clean(z["Expresion_Contextual"]))
    if m: LENS.setdefault(clean(z["Zona_ID"]), (m.group(1) or m.group(2)))

# ---- Programas publicables (excluye cuarentena) y 26 colas ----
PRG = []
for p in sheet("06_PROGRAMAS"):
    if clean(p["Documentado_Publicable"]) != "SI": continue
    PRG.append(dict(code=clean(p["Codigo"]), ar=int(p["Aridad"]), cat=clean(p["Categoria"]),
        nombre=clean(p["Nombre_Producto"]) or clean(p["Nombre_Canonico"]), trad=clean(p["Nombre_Canonico"]),
        rec=clean(p["Recurso"]), sombra=clean(p["Sombra_Bloqueo"]), tarea=clean(p["Tarea_Integracion"]),
        cola26=clean(p["Es_Cola26"]) == "SI", aviso=clean(p["Advertencia_Publicacion"])))
COL = []
for c in sheet("07_COLAS_26"):
    COL.append(dict(code=clean(c["Codigo"]), nombre=clean(c["Nombre_Producto"]) or clean(c["Nombre"]), trad=clean(c["Nombre"]),
        rec=clean(c["Recurso"]), sombra=clean(c["Sombra"]), integ=clean(c["Integracion"]), aviso=clean(c["Advertencia_Publicacion"])))
assert len(COL) == 26

# ---- Arcano anual (interpretación temporal) ----
ANU = {int(a["Arcano_ID"]): clean(a["Interpretacion_Anual"]) for a in sheet("16_ARCANO_ANUAL")
       if str(a["Arcano_ID"]).isdigit() and clean(a["Tipo"]) == "INTERPRETACION_TEMPORAL"}

# ---- .docx: diccionario Arcano × tipo de zona + núcleos de programas ----
d = docx.Document(DOCX)
KEYS = [("Retrato", "retrato"), ("Talentos", "talentos"), ("Centro", "centro"), ("Linaje paterno", "linPat"),
        ("Linaje materno", "linMat"), ("Amor", "amor"), ("Dinero", "dinero"), ("Repetición", "repeticion"),
        ("Propósito/ciclo", "proposito")]
DIC = {}
for para in d.paragraphs:
    m = re.match(r"Arcano (\d+) — [^.]+\.\s*(.*)", para.text.strip())
    if not m: continue
    body = re.sub(r"\s*Astro Lúdico usa[^.]*\.", "", m.group(2))          # notas de fuente, no contenido
    idx = [(body.find(k + ":"), k, key) for k, key in KEYS]
    idx = sorted(i for i in idx if i[0] >= 0)
    DIC[int(m.group(1))] = {key: body[pos + len(k) + 1: (idx[j + 1][0] if j + 1 < len(idx) else None)].strip()
                            for j, (pos, k, key) in enumerate(idx)}
assert len(DIC) == 22 and all(len(v) == 9 for v in DIC.values()), {k: len(v) for k, v in DIC.items()}
NUC = {}
for tb in d.tables:
    hdr = [c.text.strip() for c in tb.rows[0].cells]
    if hdr[:2] != ["Código", "Programa"]: continue
    for row in tb.rows[1:]:
        code, name, defi = (c.text.strip() for c in row.cells)
        NUC[code] = dict(nombre=name, def_=clean(defi))
ZONE_Q = {}
for tb in d.tables:
    hdr = [c.text.strip() for c in tb.rows[0].cells]
    if hdr == ["Zona", "Qué parte del programa se interpreta"]:
        for row in tb.rows[1:]: ZONE_Q[row.cells[0].text.strip()] = clean(row.cells[1].text)

# ---- v2: 770 lecturas Arcano × Posición (plantillas verificadas) ----
wv = openpyxl.load_workbook(V2, read_only=True)
def sheet2(wbk, name):
    rows = list(wbk[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if r[0] is not None]
AK = dict(Nombre_Canonico="n", Nucleo_Editorial="nuc", Necesidad_Nuclear="nec", Miedo_Nuclear="miedo",
          Semilla_Hipo="sHipo", Semilla_Hiper="sHiper", Semilla_Equilibrio="sEq", Semilla_Integracion="sInt",
          Fuente_Centro_Nucleo="cNuc", Fuente_Centro_Sombra="cSom", Fuente_Clave_Estrategica="cClave",
          Fuente_Global_Ausencia="aus", Fuente_Global_Repeticion="rep")
PK = dict(Punto="punto", Zona="zona", Posicion_Funcional="posf", Funcion="func", Pregunta_Clave="preg",
          Disparadores="disp", Hipo_Contextual="hipoC", Hiper_Contextual="hiperC", Equilibrio_Contextual="eqC",
          Prueba_Integracion="prueba")
VA = {int(r["Arcano"]): {k: clean(r[c]) for c, k in AK.items()} for r in sheet2(wv, "01_ARCANOS_BASE")}
VP = collections.OrderedDict((clean(r["Contexto_ID"]), {k: clean(r[c]) for c, k in PK.items()}) for r in sheet2(wv, "02_POSICIONES"))
assert len(VA) == 22 and len(VP) == 35
assert all(VA[a]["n"] == ARC[a]["n"] for a in VA), "nombres de arcano distintos entre Maestra y v2"
for cid, p in VP.items(): p["posl"] = p["posf"][0].lower() + p["posf"][1:]
COLS = dict(Interpretacion_Posicion="interp", Hipo="hipo", Hiper="hiper", Equilibrio="eq", Integracion="integ", Patron_Observable="pat")
TPL = {k: [] for k in COLS.values()}; USE = {k: {} for k in COLS.values()}
fill = lambda t, d: re.sub(r"\{(\w+)\}", lambda m: d[m.group(1)], t)
LECT = sheet2(wv, "03_INTERPRETACIONES_770"); assert len(LECT) == 770
for r in LECT:
    a, cid = int(r["Arcano"]), clean(r["ID_LECTURA"]).split("_", 1)[1]
    d = {**VA[a], **{k: v for k, v in VP[cid].items() if k not in ("punto", "zona")}}
    keys = sorted((k for k in d if d[k]), key=lambda k: -len(d[k]))
    for col, key in COLS.items():
        t = clean(r[col])
        for k in keys: t = t.replace(d[k], "{" + k + "}")
        if t not in TPL[key]: TPL[key].append(t)
        i = TPL[key].index(t)
        assert USE[key].setdefault(cid, i) == i, f"plantilla {key} cambia dentro de {cid}"
        assert fill(t, d) == clean(r[col]), f"no reproduce {r['ID_LECTURA']} {col}"
VAM = {a: {k: genero.masc(v) for k, v in f.items() if genero.masc(v) != v} for a, f in VA.items()}
VPM = {c: {k: genero.masc(v) for k, v in f.items() if genero.masc(v) != v} for c, f in VP.items()}
VAM = {a: f for a, f in VAM.items() if f}; VPM = {c: f for c, f in VPM.items() if f}
V2D = dict(arc=VA, pos=VP, arcM=VAM, posM=VPM, tpl=TPL, use=USE)

# ---- Tabla de programas: expresiones por zona con soporte documental específico ----
wt = openpyxl.load_workbook(TABLA, read_only=True)
SEQ_BY_PTS = {"-".join(z["pts"]): z["id"] for z in SEQ}
PUB = {p["code"] for p in PRG}
PZ = {}
for r in sheet2(wt, "03_PROGRAMA_X_ZONA"):
    if clean(r["Soporte_Zona"]) == "SINTESIS_CONTEXTUAL_MATRIX" or clean(r["Codigo"]) not in PUB: continue
    PZ[SEQ_BY_PTS[clean(r["Secuencia_Puntos"])] + "|" + clean(r["Codigo"])] = dict(txt=clean(r["Expresion_Contextual"]), fuente=clean(r["Soporte_Zona"]))
# el catálogo nuevo debe coincidir con el de la Maestra (mismos programas publicables)
TP = {clean(r["Codigo"]) for r in sheet2(wt, "01_PROGRAMAS") if clean(r["Documentado_Publicable"]) == "SI"}
assert TP == PUB, "la tabla de programas no coincide con la Maestra"

data = dict(version="MATRIX Biblioteca Maestra v1.12 + Interpretaciones v2 + diccionario semántico", arc=ARC, pts=PTS, ctx=CTX, zon=ZON,
            seq=SEQ, lens=LENS, prg=PRG, col=COL, anual=ANU, dic=DIC, nuc=NUC, zoneQ=ZONE_Q, v2=V2D, pz=PZ)
js = "/* Generado por tools/build_matrix_data.py — no editar a mano. */\nwindow.MATRIX_DATA=" + \
     json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("def_", "def") + ";\n"
open(OUT, "w", encoding="utf-8").write(js)
print(f"OK {OUT}: {len(js)/1024:.0f} KB | arcanos {len(ARC)} · puntos {len(PTS)} · contextos {len(CTX)} (770 párrafos verificados) · "
      f"zonas {len(ZON)} · secuencias {len(SEQ)} · lentes {len(LENS)} · programas publicables {len(PRG)} · colas {len(COL)} · "
      f"anual {len(ANU)} · diccionario {len(DIC)}×9 · núcleos docx {len(NUC)} · zonas docx {len(ZONE_Q)} · "
      f"v2 770 lecturas (plantillas {', '.join(f'{k}:{len(v)}' for k, v in TPL.items())}; género M en {len(VAM)} arcanos/{len(VPM)} posiciones) · "
      f"programa×zona documentados {len(PZ)}")
